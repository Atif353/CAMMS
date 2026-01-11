#!/usr/bin/env python3

import sys
import json
import base64
import io
from datetime import datetime
import warnings

warnings.filterwarnings("ignore")

import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, PieChart, Reference

THEME_COLORS = {
    "primary": "#003f5c",
    "secondary": "#2f4b7c",
    "accent": "#665191",
    "success": "#2ecc71",
    "warning": "#f1c40f",
    "danger": "#e74c3c",
    "light_bg": "#f7f9fb",
    "grid": "#e0e0e0",
}

plt.rcParams.update({
    "font.family": "DejaVu Sans",
    "axes.titlesize": 14,
    "axes.labelsize": 11,
    "xtick.labelsize": 9,
    "ytick.labelsize": 9,
    "figure.titlesize": 16,
    "axes.grid": True,
    "grid.alpha": 0.3,
    "axes.axisbelow": True,
})

def build_overall_summary(report_data: dict) -> dict:
    instructors = report_data.get("instructorSummary", [])
    total_instructors = len(instructors)
    total_time = 0
    total_completed = 0
    total_pending = 0
    total_activities = 0
    activity_type_aggregate = {}

    for instr in instructors:
        total_time += instr.get("overallTimeSpent", 0) or 0
        total_completed += instr.get("completedActivities", 0) or 0
        total_pending += instr.get("pendingActivities", 0) or 0
        totals = instr.get("totals", {}) or {}

        for at_code, stats in totals.items():
            count = stats.get("count", 0) or 0
            time_spent = stats.get("timeSpent", 0) or 0
            total_activities += count

            if at_code not in activity_type_aggregate:
                activity_type_aggregate[at_code] = {"count": 0, "timeSpent": 0}

            activity_type_aggregate[at_code]["count"] += count
            activity_type_aggregate[at_code]["timeSpent"] += time_spent

    completion_rate = 0
    if total_activities > 0:
        completion_rate = round((total_completed / total_activities) * 100, 2)

    return {
        "totalInstructors": total_instructors,
        "totalActivities": total_activities,
        "totalTimeSpent": total_time,
        "totalCompleted": total_completed,
        "totalPending": total_pending,
        "completionRate": completion_rate,
        "activityTypeAggregate": activity_type_aggregate,
    }

def format_time(minutes):
    if minutes < 60:
        return f"{int(minutes)} min"
    hours = minutes // 60
    mins = int(minutes % 60)
    if mins == 0:
        return f"{int(hours)} hr"
    return f"{int(hours)} hr {mins} min"

def add_title_page(pdf: PdfPages, report_data: dict, report_type: str):
    fig = plt.figure(figsize=(11.7, 8.3))
    ax = fig.add_subplot(111)
    ax.set_facecolor(THEME_COLORS["primary"])
    ax.axis("off")

    start = report_data.get("startDate", "")
    end = report_data.get("endDate", "")
    generated_at = report_data.get("generatedAt") or datetime.utcnow().isoformat()

    fig.text(
        0.5, 0.75,
        "Course Activity Management & Monitoring System",
        ha="center", va="center",
        fontsize=24, color="white", weight="bold"
    )

    fig.text(
        0.5, 0.65,
        f"{report_type.title()} Activity Report",
        ha="center", va="center",
        fontsize=18, color="#ecf0f1"
    )

    fig.text(
        0.5, 0.55,
        f"Period: {start}  →  {end}",
        ha="center", va="center",
        fontsize=12, color="#dcdde1"
    )

    fig.text(
        0.5, 0.47,
        f"Generated on: {generated_at}",
        ha="center", va="center",
        fontsize=10, color="#bdc3c7"
    )

    fig.text(
        0.01, 0.02,
        "Powered by CAMMS",
        ha="left", va="center",
        fontsize=8, color="#ecf0f1"
    )

    pdf.savefig(fig, bbox_inches="tight")
    plt.close(fig)

def add_overall_summary_page(pdf: PdfPages, report_data: dict):
    overall = build_overall_summary(report_data)
    instructors = report_data.get("instructorSummary", [])

    fig, axes = plt.subplots(1, 2, figsize=(11.7, 8.3))
    fig.suptitle("Executive Summary", fontsize=16, fontweight="bold", color=THEME_COLORS["primary"])

    ax = axes[0]
    ax.axis("off")
    ax.set_facecolor(THEME_COLORS["light_bg"])

    lines = [
        f"Total Instructors: {overall['totalInstructors']}",
        f"Total Activities: {overall['totalActivities']}",
        f"Total Time Spent: {format_time(overall['totalTimeSpent'])}",
        f"Completed Activities: {overall['totalCompleted']}",
        f"Pending Activities: {overall['totalPending']}",
        f"Completion Rate: {overall['completionRate']}%",
    ]

    y = 0.9
    for text in lines:
        ax.text(0.05, y, text, fontsize=11, color=THEME_COLORS["secondary"], weight="bold")
        y -= 0.08

    ax2 = axes[1]
    agg = overall["activityTypeAggregate"]

    if agg:
        labels = list(agg.keys())
        counts = [agg[a]["count"] for a in labels]

        max_label_len = 20
        abbreviated_labels = [label[:max_label_len] + "..." if len(label) > max_label_len else label for label in labels]

        colors = plt.cm.Set3(range(len(labels)))
        bars = ax2.bar(abbreviated_labels, counts, color=colors, edgecolor=THEME_COLORS["primary"], linewidth=1.5)

        ax2.set_title("Activities by Type", fontsize=12, fontweight="bold", pad=15)
        ax2.set_xlabel("Activity Type", fontsize=10, fontweight="bold")
        ax2.set_ylabel("Count", fontsize=10, fontweight="bold")
        ax2.tick_params(axis="x", rotation=45, labelsize=8)
        ax2.grid(axis="y", alpha=0.3, linestyle="--", color=THEME_COLORS["grid"])

        max_count = max(counts) if counts else 0
        padding = max(max_count * 0.15, 5)
        ax2.set_ylim(0, max_count + padding)
        
        for bar, v in zip(bars, counts):
            height = bar.get_height()
            label_y = height + padding * 0.05
            ax2.text(bar.get_x() + bar.get_width() / 2., label_y,
                    str(v), ha="center", va="bottom", fontsize=9, fontweight="bold",
                    bbox=dict(boxstyle="round,pad=0.3", facecolor="white", edgecolor=THEME_COLORS["primary"], alpha=0.8))

        plt.tight_layout(rect=[0, 0, 1, 0.95])
    else:
        ax2.axis("off")
        ax2.text(
            0.5, 0.5,
            "No activity data available for this period.",
            ha="center", va="center",
            fontsize=11, color="gray"
        )

    plt.tight_layout(rect=[0, 0, 1, 0.95])
    pdf.savefig(fig, bbox_inches="tight")
    plt.close(fig)

    if len(instructors) > 1:
        fig, ax = plt.subplots(figsize=(11.7, 8.3))
        names = [i.get("name", "Unknown")[:25] + "..." if len(i.get("name", "Unknown")) > 25 else i.get("name", "Unknown") for i in instructors]
        times = [i.get("overallTimeSpent", 0) for i in instructors]

        colors = plt.cm.viridis([i / max(len(times), 1) for i in range(len(times))])
        bars = ax.barh(names, times, color=colors, edgecolor=THEME_COLORS["primary"], linewidth=1.5)

        ax.set_title("Time Spent by Instructor", fontsize=14, fontweight="bold", pad=15)
        ax.set_xlabel("Time (Minutes)", fontsize=11, fontweight="bold")
        ax.set_ylabel("Instructor", fontsize=11, fontweight="bold")
        ax.invert_yaxis()
        ax.grid(axis="x", alpha=0.3, linestyle="--", color=THEME_COLORS["grid"])

        max_time = max(times) if times else 0
        padding = max(max_time * 0.15, 100)
        ax.set_xlim(0, max_time + padding)

        for i, (bar, v) in enumerate(zip(bars, times)):
            label_text = format_time(v)
            label_x = v + padding * 0.1
            ax.text(label_x, bar.get_y() + bar.get_height() / 2,
                   label_text, va="center", fontsize=9, fontweight="bold", 
                   bbox=dict(boxstyle="round,pad=0.3", facecolor="white", edgecolor=THEME_COLORS["primary"], alpha=0.8))

        plt.tight_layout()
        pdf.savefig(fig, bbox_inches="tight")
        plt.close(fig)

def add_instructor_pages(pdf: PdfPages, report_data: dict):
    instructors = report_data.get("instructorSummary", [])

    if not instructors:
        return

    for instr in instructors:
        name = instr.get("name", "Unknown Instructor")
        email = instr.get("email", "N/A")
        totals = instr.get("totals", {}) or {}

        fig, axes = plt.subplots(2, 1, figsize=(11.7, 8.3), height_ratios=[1, 2])
        fig.suptitle(f"Instructor: {name}", fontsize=14, fontweight="bold", color=THEME_COLORS["primary"])

        ax_info = axes[0]
        ax_info.set_facecolor(THEME_COLORS["light_bg"])
        ax_info.axis("off")

        header_lines = [
            f"Name: {name}",
            f"Email: {email}",
            f"Total Time Spent: {format_time(instr.get('overallTimeSpent', 0))}",
            f"Completed Activities: {instr.get('completedActivities', 0)}",
            f"Pending Activities: {instr.get('pendingActivities', 0)}",
        ]

        y = 0.85
        for text in header_lines:
            ax_info.text(0.05, y, text, fontsize=10, color=THEME_COLORS["secondary"], weight="bold")
            y -= 0.15

        if totals:
            ax_chart = axes[1]
            activity_types = list(totals.keys())
            time_spent = [totals[a]["timeSpent"] for a in activity_types]

            max_label_len = 20
            abbreviated_types = [at[:max_label_len] + "..." if len(at) > max_label_len else at for at in activity_types]

            colors = plt.cm.Pastel1(range(len(activity_types)))
            bars = ax_chart.bar(abbreviated_types, time_spent, color=colors, edgecolor=THEME_COLORS["primary"], linewidth=1.5)

            ax_chart.set_title("Time Spent by Activity Type", fontsize=12, fontweight="bold", pad=15)
            ax_chart.set_xlabel("Activity Type", fontsize=10, fontweight="bold")
            ax_chart.set_ylabel("Time (Minutes)", fontsize=10, fontweight="bold")
            ax_chart.tick_params(axis="x", rotation=45, labelsize=8)
            ax_chart.grid(axis="y", alpha=0.3, linestyle="--", color=THEME_COLORS["grid"])

            max_time = max(time_spent) if time_spent else 0
            padding = max(max_time * 0.15, 50)
            ax_chart.set_ylim(0, max_time + padding)
            
            for bar, v in zip(bars, time_spent):
                height = bar.get_height()
                label_text = format_time(v)
                label_y = height + padding * 0.05
                ax_chart.text(bar.get_x() + bar.get_width() / 2., label_y,
                        label_text, ha="center", va="bottom", fontsize=9, fontweight="bold",
                        bbox=dict(boxstyle="round,pad=0.3", facecolor="white", edgecolor=THEME_COLORS["primary"], alpha=0.8))

        plt.tight_layout(rect=[0, 0, 1, 0.95])
        pdf.savefig(fig, bbox_inches="tight")
        plt.close(fig)

        if totals:
            fig2, ax2 = plt.subplots(figsize=(11.7, 8.3))
            ax2.axis("off")
            ax2.set_facecolor(THEME_COLORS["light_bg"])

            table_data = []
            for at, stats in totals.items():
                count = stats.get("count", 0) or 0
                t_spent = stats.get("timeSpent", 0) or 0
                avg_time = round(t_spent / count, 2) if count > 0 else 0
                table_data.append([at, count, format_time(t_spent), format_time(avg_time)])

            col_labels = ["Activity Type", "Count", "Total Time", "Avg Time per Activity"]
            table = ax2.table(
                cellText=table_data,
                colLabels=col_labels,
                cellLoc="center",
                loc="center",
                colWidths=[0.4, 0.15, 0.2, 0.25]
            )

            table.auto_set_font_size(False)
            table.set_fontsize(9)
            table.scale(1, 2)

            for i in range(len(col_labels)):
                cell = table[(0, i)]
                cell.set_facecolor(THEME_COLORS["primary"])
                cell.set_text_props(weight="bold", color="white")

            for i in range(1, len(table_data) + 1):
                for j in range(len(col_labels)):
                    cell = table[(i, j)]
                    if i % 2 == 0:
                        cell.set_facecolor(THEME_COLORS["light_bg"])
                    cell.set_edgecolor(THEME_COLORS["grid"])

            ax2.set_title("Activity Breakdown", fontsize=12, fontweight="bold", pad=20)
            pdf.savefig(fig2, bbox_inches="tight")
            plt.close(fig2)

def generate_pdf_report(report_data: dict, report_type: str) -> bytes:
    buffer = io.BytesIO()
    with PdfPages(buffer) as pdf:
        add_title_page(pdf, report_data, report_type)
        add_overall_summary_page(pdf, report_data)
        add_instructor_pages(pdf, report_data)
    buffer.seek(0)
    return buffer.getvalue()

def style_worksheet(ws):
    header_fill = PatternFill(start_color="003f5c", end_color="003f5c", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    center_alignment = Alignment(horizontal="center", vertical="center")
    left_alignment = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )

    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_alignment
        cell.border = thin_border

    for row in range(2, ws.max_row + 1):
        fill = None
        if row % 2 == 0:
            fill = PatternFill(start_color="f7f9fb", end_color="f7f9fb", fill_type="solid")

        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=row, column=col)
            if fill:
                cell.fill = fill
            cell.border = thin_border
            if col == 1:
                cell.alignment = left_alignment
            else:
                cell.alignment = center_alignment

    for col in range(1, ws.max_column + 1):
        column_letter = get_column_letter(col)
        max_length = 0
        for row in range(1, ws.max_row + 1):
            cell = ws[f"{column_letter}{row}"]
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[column_letter].width = max(12, min(max_length + 2, 50))

    ws.freeze_panes = "A2"

def add_chart_to_worksheet(ws, chart_type, data_col, title, position):
    try:
        if ws.max_row < 2:
            return

        if chart_type == "bar":
            chart = BarChart()
            chart.type = "col"
            chart.style = 10
            chart.title = title
            chart.y_axis.title = "Value"
            chart.x_axis.title = "Category"
            data = Reference(ws, min_col=data_col, min_row=1, max_row=ws.max_row)
            cats = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)

        elif chart_type == "pie":
            chart = PieChart()
            chart.title = title
            data = Reference(ws, min_col=data_col, min_row=1, max_row=ws.max_row)
            cats = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)

        chart.width = 12
        chart.height = 8
        ws.add_chart(chart, position)
    except Exception:
        pass

def generate_excel_report(report_data: dict, report_type: str) -> bytes:
    buffer = io.BytesIO()
    instructors = report_data.get("instructorSummary", [])
    overall = build_overall_summary(report_data)

    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        summary_rows = []
        for instr in instructors:
            total_activities = 0
            for stats in (instr.get("totals") or {}).values():
                total_activities += stats.get("count", 0) or 0

            completion_rate = 0
            if total_activities > 0:
                completion_rate = round(
                    (instr.get("completedActivities", 0) or 0) / total_activities * 100, 2
                )

            summary_rows.append({
                "Instructor Name": instr.get("name", "Unknown"),
                "Email": instr.get("email", ""),
                "Total Activities": total_activities,
                "Total Time Spent (min)": instr.get("overallTimeSpent", 0),
                "Completed Activities": instr.get("completedActivities", 0),
                "Pending Activities": instr.get("pendingActivities", 0),
                "Completion Rate (%)": completion_rate,
            })

        df_summary = pd.DataFrame(summary_rows) if summary_rows else pd.DataFrame(
            columns=[
                "Instructor Name", "Email", "Total Activities",
                "Total Time Spent (min)", "Completed Activities",
                "Pending Activities", "Completion Rate (%)",
            ]
        )

        df_summary.to_excel(writer, sheet_name="Summary", index=False)
        ws_summary = writer.book["Summary"]
        style_worksheet(ws_summary)

        if len(summary_rows) > 0:
            add_chart_to_worksheet(
                ws_summary,
                "bar",
                4,
                "Time Spent by Instructor",
                "I2"
            )

        last_row = ws_summary.max_row + 2
        ws_summary[f"A{last_row}"] = "Overall Totals"
        ws_summary[f"A{last_row}"].font = Font(bold=True, size=12, color="003f5c")
        ws_summary[f"A{last_row}"].fill = PatternFill(start_color="f7f9fb", end_color="f7f9fb", fill_type="solid")

        ws_summary[f"A{last_row + 1}"] = "Total Instructors"
        ws_summary[f"B{last_row + 1}"] = overall["totalInstructors"]
        ws_summary[f"A{last_row + 2}"] = "Total Activities"
        ws_summary[f"B{last_row + 2}"] = overall["totalActivities"]
        ws_summary[f"A{last_row + 3}"] = "Total Time Spent (min)"
        ws_summary[f"B{last_row + 3}"] = overall["totalTimeSpent"]
        ws_summary[f"A{last_row + 4}"] = "Completion Rate (%)"
        ws_summary[f"B{last_row + 4}"] = overall["completionRate"]

        for row in range(last_row, last_row + 5):
            for col in range(1, 3):
                cell = ws_summary.cell(row=row, column=col)
                if row == last_row:
                    cell.border = Border(
                        left=Side(style="thin"), right=Side(style="thin"),
                        top=Side(style="thin"), bottom=Side(style="thin")
                    )
                else:
                    cell.border = Border(
                        left=Side(style="thin"), right=Side(style="thin"),
                        bottom=Side(style="thin")
                    )

        activity_type_data = []
        for at_code, stats in overall["activityTypeAggregate"].items():
            activity_type_data.append({
                "Activity Type": at_code,
                "Count": stats["count"],
                "Total Time (min)": stats["timeSpent"],
            })

        if activity_type_data:
            df_activity = pd.DataFrame(activity_type_data)
            df_activity.to_excel(writer, sheet_name="Activity Types", index=False)
            ws_activity = writer.book["Activity Types"]
            style_worksheet(ws_activity)

            try:
                add_chart_to_worksheet(
                    ws_activity,
                    "pie",
                    None,
                    None,
                    "Activities Distribution",
                    "E2"
                )
            except:
                pass

        for instr in instructors:
            totals = instr.get("totals", {}) or {}
            if not totals:
                continue

            detail_data = []
            for at_code, stats in totals.items():
                count = stats.get("count", 0) or 0
                time_spent = stats.get("timeSpent", 0) or 0
                avg_time = round(time_spent / count, 2) if count > 0 else 0

                detail_data.append({
                    "Activity Type": at_code,
                    "Count": count,
                    "Time Spent (min)": time_spent,
                    "Avg Time per Activity (min)": avg_time,
                })

            df_detail = pd.DataFrame(detail_data)
            sheet_name = (instr.get("name", "Instructor")[:31] or "Instructor").replace("/", "-").replace("\\", "-")
            df_detail.to_excel(writer, sheet_name=sheet_name, index=False)
            ws = writer.book[sheet_name]
            style_worksheet(ws)

            add_chart_to_worksheet(
                ws,
                "bar",
                3,
                "Time Spent by Activity Type",
                "F2"
            )

    buffer.seek(0)
    return buffer.getvalue()

def main():
    try:
        input_data = json.load(sys.stdin)
        report_data = input_data.get("reportData", {})
        report_type = input_data.get("reportType", "WEEKLY")
        output_format = input_data.get("outputFormat", "pdf").lower()

        if output_format == "pdf":
            file_bytes = generate_pdf_report(report_data, report_type)
        elif output_format == "excel":
            file_bytes = generate_excel_report(report_data, report_type)
        else:
            raise ValueError(f"Unsupported output format: {output_format}")

        file_base64 = base64.b64encode(file_bytes).decode("utf-8")
        result = {
            "success": True,
            "fileData": file_base64,
            "format": output_format,
        }
        print(json.dumps(result))

    except Exception as e:
        error_result = {
            "success": False,
            "error": str(e),
        }
        print(json.dumps(error_result))
        sys.exit(1)

if __name__ == "__main__":
    main()
